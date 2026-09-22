# -*- coding: utf-8 -*-
"""Build the JBM DMS Service Mobile App user manual as a .docx, laying every
screen out in the fig-1 style: blue step heading, square bullets, a blue band
and a fully annotated screenshot."""
import os, sys, json
sys.path.insert(0, "lib")
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import shapes
import spec_a, spec_b, spec_c

LAYOUTS = json.load(open("figs/layouts.json"))

STEPS = spec_a.STEPS + spec_b.STEPS + spec_c.STEPS

BLUE      = RGBColor(0x1C, 0xA2, 0xDB)     # fig-1 heading blue
BAND      = "4E81BD"                        # fig-1 band fill
DARK      = RGBColor(0x20, 0x20, 0x20)
GREY      = RGBColor(0x60, 0x60, 0x60)
NAVY      = RGBColor(0x0F, 0x3D, 0x66)
FIG_W_IN  = 8.27 - 2 * 0.62        # full text column on A4
FIG_W     = Inches(FIG_W_IN)
FONT      = "Arial"

ROMAN = [(1000,"m"),(900,"cm"),(500,"d"),(400,"cd"),(100,"c"),(90,"xc"),
         (50,"l"),(40,"xl"),(10,"x"),(9,"ix"),(5,"v"),(4,"iv"),(1,"i")]


def roman(n):
    out = ""
    for v, s in ROMAN:
        while n >= v:
            out += s; n -= v
    return out


# ---------------------------------------------------------------- helpers
def shade(el, fill):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), fill)
    el.append(sh)


def border(par, edge="bottom", sz=12, color="000000"):
    pPr = par._p.get_or_add_pPr()
    bd = pPr.find(qn("w:pBdr"))
    if bd is None:
        bd = OxmlElement("w:pBdr"); pPr.append(bd)
    e = OxmlElement(f"w:{edge}")
    e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), "1"); e.set(qn("w:color"), color)
    bd.append(e)


def run(par, text, size=11, bold=False, color=DARK, italic=False, font=FONT):
    r = par.add_run(text)
    r.font.name = font; r.font.size = Pt(size); r.bold = bold; r.italic = italic
    r.font.color.rgb = color
    r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    return r


def para(doc, space_before=0, space_after=4, align=None, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.0
    if align is not None:
        p.alignment = align
    if keep:
        p.paragraph_format.keep_with_next = True
    return p


def field(par, instr):
    for kind, val in (("begin", None), ("instrText", instr), ("end", None)):
        r = OxmlElement("w:r")
        if kind == "instrText":
            t = OxmlElement("w:instrText"); t.set(qn("xml:space"), "preserve"); t.text = val
            r.append(t)
        else:
            f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), kind); r.append(f)
        par._p.append(r)


# ---------------------------------------------------------------- blocks
def rule(doc, before=0, after=6):
    p = para(doc, before, after, keep=True)
    border(p, "bottom", sz=12)
    p.paragraph_format.space_after = Pt(after)
    return p


def step_heading(doc, num, title):
    p = para(doc, 0, 6, keep=True)
    run(p, f"{roman(num)}. ", size=19, bold=True, color=BLUE)
    run(p, title, size=19, bold=True, color=BLUE)
    return p


def bullet(doc, text):
    p = para(doc, 0, 3, keep=True)
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.22)
    run(p, "▪  ", size=12, bold=True)
    run(p, text, size=12, bold=True)
    return p


def band(doc, text):
    p = para(doc, 6, 8, keep=True)
    shade(p._p.get_or_add_pPr(), BAND)
    p.paragraph_format.left_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(7)
    run(p, text, size=10.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    return p


def figure(doc, path, caption, lay, width_in=FIG_W_IN):
    """Place the bare device shot, then lay the callouts over it as Word shapes.

    Everything a reader may want to change - the callout wording, the size of a
    box, whether an arrow is there at all - stays a separate, editable shape."""
    CW, _CH = lay["canvas"]
    PW, _PH = lay["phone"]
    scale = width_in / CW * shapes.EMU      # canvas pixels -> EMU
    p = para(doc, 0, 2, align=WD_ALIGN_PARAGRAPH.CENTER, keep=True)
    p.add_run().add_picture(path, width=Inches(PW * width_in / CW))
    for it in lay["items"]:
        tx, ty, tw, th = it["target"]
        ax, ay, px, py = it["arrow"]
        bx, by, bw, bh = it["box"]
        # a little slack, so Word never re-wraps the text the layout measured;
        # the box grows away from the arrow, keeping the leader on its edge
        grow = bw * 0.07
        if it["side"] == "L":
            bx -= grow
        bw += grow
        shapes.outline(p, tx * scale, ty * scale, tw * scale, th * scale, line_pt=1.0)
        shapes.arrow(p, ax * scale, ay * scale, px * scale, py * scale, line_pt=0.75)
        shapes.textbox(p, bx * scale, by * scale, bw * scale, (bh + 10) * scale,
                       it["text"], pt=7.5, font=FONT, line_pt=1.0)
    c = para(doc, 0, 6, align=WD_ALIGN_PARAGRAPH.CENTER)
    run(c, caption, size=9, italic=True, color=GREY)
    return p


def note(doc, text):
    p = para(doc, 3, 4)
    shade(p._p.get_or_add_pPr(), "EDF2F8")
    p.paragraph_format.left_indent = Inches(0.08)
    run(p, "Note:  ", size=9.5, bold=True, color=NAVY)
    run(p, text, size=9.5, color=NAVY)
    return p


def h1(doc, text, pagebreak=True):
    if pagebreak:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    p = para(doc, 0, 4, keep=True)
    border(p, "bottom", sz=12)
    run(p, text, size=22, bold=True, color=BLUE)
    return p


def h2(doc, text):
    p = para(doc, 10, 4, keep=True)
    run(p, text, size=13, bold=True, color=NAVY)
    return p


def body(doc, text, size=11):
    p = para(doc, 0, 5)
    run(p, text, size=size)
    return p


def dash(doc, text):
    p = para(doc, 0, 3)
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.22)
    run(p, "▪  ", size=11)
    run(p, text, size=11)
    return p


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htxt in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        shade(cell._tc.get_or_add_tcPr(), BAND)
        p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
        run(p, htxt, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ""
            p = cells[i].paragraphs[0]; p.paragraph_format.space_after = Pt(2)
            run(p, v, size=10)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    para(doc, 2, 8)
    return t


# ---------------------------------------------------------------- content
CH_OBJECTIVE = {
 "2. Login / Logout":
   "To allow authorized JBM Service and JBM Vendor depot users to securely access the "
   "JBM DMS Mobile Application.",
 "3. User Profile":
   "To allow users to view their profile and update the parameters permitted by JBM.",
 "4. Contact Us":
   "To provide depot users with JBM HQ contact information for support and communication.",
 "5. Gate In / Gate Out":
   "The Gate In / Gate Out module is used to record the arrival and release of depot vehicles, "
   "capture driver complaints, and support the service process before a vehicle is released for "
   "out-shedding.",
 "6. Auto Job Cards (Daily / Ten-Day)":
   "To provide depot users with the list of auto-generated jobs based on daily and ten-day "
   "service schedules.",
 "7. Assign Auto Job Cards":
   "To assign auto generated jobs to the appropriate service supervisor or authorized users.",
 "8. DCR Charging":
   "To allow depot users to view and capture vehicle charging-related information, as enabled "
   "by the JBM DMS application.",
}

NAV_ICONS = [
 ("Login / Logout",        "docimg/png02.png"),
 ("Menu Bar",              "docimg/png03.png"),
 ("User Profile",          "docimg/png04.png"),
 ("Contact Us",            "docimg/png05.png"),
 ("Gate In / Gate Out",    "docimg/png06.png"),
 ("Auto Job Card List",    "docimg/png07.png"),
 ("Assign Auto Jobs",      "docimg/png08.png"),
 ("DCR Charging",          "docimg/png09.png"),
 ("Dashboard",             "docimg/png10.png"),
 ("Scroll Right",          "docimg/png11.png"),
 ("Scroll Left",           "docimg/png12.png"),
]


def page_setup(doc):
    s = doc.sections[0]
    s.page_width, s.page_height = Inches(8.27), Inches(11.69)   # A4 portrait
    s.left_margin = s.right_margin = Inches(0.62)
    s.top_margin = Inches(0.55); s.bottom_margin = Inches(0.5)
    s.header_distance = Inches(0.28); s.footer_distance = Inches(0.25)

    hp = s.header.paragraphs[0]
    hp.paragraph_format.space_after = Pt(2)
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if os.path.exists("docimg/png01.png"):
        hp.add_run().add_picture("docimg/png01.png", width=Inches(0.95))
    hp.add_run("   ")
    run(hp, "Dealer Management System  ·  User Manual  ·  Service Mobile App",
        size=8.5, color=GREY)
    border(hp, "bottom", sz=6, color="BFBFBF")

    fp = s.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    border(fp, "top", sz=6, color="BFBFBF")
    run(fp, "HRSL Confidential", size=8, color=GREY)
    fp.add_run("\t\t")
    run(fp, "Page ", size=8, color=GREY)
    field(fp, " PAGE ")
    run(fp, " of ", size=8, color=GREY)
    field(fp, " NUMPAGES ")


def cover(doc):
    para(doc, 40, 0)
    p = para(doc, 0, 6, align=WD_ALIGN_PARAGRAPH.CENTER)
    if os.path.exists("jbm_logo.png"):
        p.add_run().add_picture("jbm_logo.png", width=Inches(2.7))
    para(doc, 18, 0)
    p = para(doc, 0, 2, align=WD_ALIGN_PARAGRAPH.CENTER)
    run(p, "Dealer Management System", size=28, bold=True, color=BLUE)
    p = para(doc, 0, 2, align=WD_ALIGN_PARAGRAPH.CENTER)
    run(p, "User Manual", size=20, bold=True, color=NAVY)
    p = para(doc, 0, 16, align=WD_ALIGN_PARAGRAPH.CENTER)
    run(p, "Service — Mobile App", size=16, color=GREY)
    rule(doc, 6, 10)
    p = para(doc, 0, 2, align=WD_ALIGN_PARAGRAPH.CENTER)
    run(p, "Document:  USER MANUAL-MOBILE APP_V1", size=11, bold=True)
    p = para(doc, 0, 2, align=WD_ALIGN_PARAGRAPH.CENTER)
    run(p, "Release Date:  10th Sept, 2026", size=11, bold=True)
    para(doc, 20, 0)
    p = para(doc, 0, 10, align=WD_ALIGN_PARAGRAPH.CENTER)
    run(p, "Every application screen in this manual is annotated field by field, "
           "in the style of the JBM DMS screen guides.", size=10.5, italic=True, color=GREY)
    para(doc, 24, 0)
    p = para(doc, 0, 4, align=WD_ALIGN_PARAGRAPH.CENTER)
    run(p, "Prepared by", size=9, color=GREY)
    p = para(doc, 0, 0, align=WD_ALIGN_PARAGRAPH.CENTER)
    if os.path.exists("docimg/png01.png"):
        p.add_run().add_picture("docimg/png01.png", width=Inches(1.9))


def front_matter(doc):
    h1(doc, "Document Control")
    h2(doc, "Document History")
    table(doc, ["Version", "Date of Prepare", "Prepared / Updated By", "Reviewed By",
                "Reason for Change", "Affected Sections"],
          [["1.0", "10th Sept 2026", "Harbans Kumar", "", "Initial release", "All"]],
          [0.7, 1.2, 1.5, 1.1, 1.4, 1.1])
    h2(doc, "Revision History")
    table(doc, ["Version", "Date of Revision", "Prepared / Updated By", "Reviewed By",
                "Reason for Change", "Affected Sections"],
          [["", "", "", "", "", ""]], [0.7, 1.2, 1.5, 1.1, 1.4, 1.1])
    h2(doc, "Approvals")
    table(doc, ["Version", "Date of Revision", "Prepared / Updated By", "Reviewed By",
                "Reason for Change", "Affected Sections"],
          [["", "", "", "", "", ""]], [0.7, 1.2, 1.5, 1.1, 1.4, 1.1])

    h1(doc, "Contents")
    chapters = [
      ("1.", "Introduction", "1.1 Purpose · 1.2 Intended Users · 1.3 Prerequisite · 1.4 General Navigation"),
    ]
    seen, order = {}, []
    for st in STEPS:
        seen.setdefault(st["ch"], []).append(st)
        if st["ch"] not in order:
            order.append(st["ch"])
    n = 0
    for ch in order:
        items = []
        for st in seen[ch]:
            n += 1
            items.append(f"{roman(n)}. {st['title']}")
        num, name = ch.split(". ", 1)
        chapters.append((num + ".", name, "  ·  ".join(items)))
    chapters.append(("9.", "User Troubleshooting", ""))
    chapters.append(("10.", "JBM Officials Interviewed", ""))
    for num, name, detail in chapters:
        p = para(doc, 6, 1)
        run(p, f"{num}  {name}", size=12, bold=True, color=NAVY)
        if detail:
            d = para(doc, 0, 4)
            d.paragraph_format.left_indent = Inches(0.3)
            run(d, detail, size=9.5, color=GREY)


def introduction(doc):
    h1(doc, "1.  Introduction")
    h2(doc, "1.1  Purpose")
    body(doc, "The purpose of this manual is to provide step-by-step guidance to depot users for "
              "performing service-related activities through the JBM DMS Mobile Application.")
    body(doc, "The application enables depot users to:")
    for t in ["Access the application using authorized credentials.",
              "View and update permitted user profile information.",
              "Access JBM HQ contact details.",
              "Capture vehicle Gate-In and Gate-Out information.",
              "Identify vehicles using QR code, VIN, Bus Number, or Registration Number.",
              "Record driver complaints, breakdown and accidental information while gating in, "
              "along with other details.",
              "Trigger and manage auto-generated service jobs.",
              "Perform charging, vendor job card, and vehicle history activities, as enabled for the user."]:
        dash(doc, t)

    h2(doc, "1.2  Intended Users")
    table(doc, ["User", "Access"],
          [["JBM Service Depot User", "Authorized service operations"],
           ["JBM Vendor Depot User", "Authorized vendor operations"]], [3.0, 4.0])

    h2(doc, "1.3  Prerequisite")
    body(doc, "The depot vehicle inventory must be available in the system before performing "
              "vehicle-related transactions. The inventory should contain the vehicle information "
              "required to identify and process depot vehicles, such as:")
    for t in ["Vehicle Identification Number", "Vehicle Registration Number", "Bus Number",
              "Assigned depot / vehicle mapping"]:
        dash(doc, t)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    h2(doc, "1.4  General Navigation")
    body(doc, "The mobile application provides the following modules through the main menu or "
              "dashboard. The exact screen arrangement and labels are subject to the approved "
              "application design. Vendor Job Card, Vehicle History and Shift Remarks are opened "
              "from the same menu and dashboard, and are available where enabled for the user.")
    t = doc.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    for i, htxt in enumerate(["Functionality", "Symbol / Icon"]):
        cell = t.rows[0].cells[i]; cell.text = ""
        shade(cell._tc.get_or_add_tcPr(), BAND)
        p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
        run(p, htxt, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    for name, icon in NAV_ICONS:
        cells = t.add_row().cells
        cells[0].text = ""
        p = cells[0].paragraphs[0]; p.paragraph_format.space_after = Pt(3)
        run(p, name, size=10)
        cells[1].text = ""
        p = cells[1].paragraphs[0]
        p.paragraph_format.space_after = Pt(3)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if os.path.exists(icon):
            p.add_run().add_picture(icon, height=Inches(0.24))
    for row in t.rows:
        row.cells[0].width = Inches(4.6); row.cells[1].width = Inches(2.4)
    para(doc, 2, 6)


def steps(doc):
    ch_seen, order = {}, []
    for st in STEPS:
        ch_seen.setdefault(st["ch"], []).append(st)
        if st["ch"] not in order:
            order.append(st["ch"])

    n = 0
    for ch in order:
        h1(doc, ch.replace(". ", ".  ", 1))
        h2(doc, "Objective")
        body(doc, CH_OBJECTIVE[ch])
        total = len(ch_seen[ch])
        h2(doc, "In this section")
        base = n
        for j, s_ in enumerate(ch_seen[ch], 1):
            q = para(doc, 0, 3)
            q.paragraph_format.left_indent = Inches(0.28)
            q.paragraph_format.first_line_indent = Inches(-0.28)
            run(q, f"{roman(base + j)}.  ", size=11, bold=True, color=BLUE)
            run(q, s_["title"], size=11, bold=True)
            c = para(doc, 0, 5)
            c.paragraph_format.left_indent = Inches(0.28)
            run(c, s_["caption"], size=9.5, italic=True, color=GREY)
        if ch.startswith("2."):
            login_validations(doc)
        for k, st in enumerate(ch_seen[ch], 1):
            n += 1
            i = STEPS.index(st) + 1
            fig = f"figs/fig{i:02d}_{st['screen']}.png"
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            rule(doc, 0, 5)
            step_heading(doc, n, st["title"])
            for b in st["bullets"]:
                bullet(doc, b)
            band(doc, f"{ch.split('. ', 1)[1].upper()}   ·   STEP {k} OF {total}")
            figure(doc, fig, f"Fig. {n} — {st['caption']}",
                   LAYOUTS[f"fig{i:02d}_{st['screen']}"])
            if st.get("note"):
                note(doc, st["note"])


def back_matter(doc):
    h1(doc, "9.  User Troubleshooting")
    table(doc, ["Issue", "Recommended action"],
      [["Unable to login", "Verify credentials and contact administrator if access is denied."],
       ["Vehicle not found", "Verify Bus Number, Registration Number, or VIN."],
       ["VIN not authorized", "Contact admin for assigned VIN access."],
       ["Odometer validation error", "Enter a reading not less than the previous recorded reading."],
       ["Duplicate Gate-In error", "Complete the previous Gate-Out loop first."],
       ["Gate-Out not allowed", "Verify Gate-In status and job card closure."],
       ["Wheel temperature alert", "Inform Shift Supervisor / Depot Manager and follow the service process."],
       ["Complaint not visible", "Verify the Gate-In complaint was saved against the correct VIN."],
       ["Job card not available", "Verify auto-job generation and assigned access."]],
      [2.4, 4.6])

    h1(doc, "10.  JBM Officials Interviewed")
    table(doc, ["Persons interviewed", "Department", "Remarks"],
      [["Mr. Puneet Sharma", "", ""], ["Mr. Gaurav Singla", "", ""],
       ["Mr. Dinesh Kumar", "", ""], ["Mr. Vinit Chavan", "", ""]], [2.6, 2.2, 2.2])


def login_validations(doc):
    h2(doc, "Login validations")
    table(doc, ["Validation", "Expected behaviour"],
      [["Valid credentials", "User can access the application"],
       ["Invalid username / password", "Login is rejected"],
       ["Unauthorized user", "Access is not permitted"],
       ["Blank mandatory fields", "System prompts the user to enter required details"]],
      [2.6, 4.4])


def renumber_drawings(doc):
    """Give every drawing in a part a unique id; python-docx numbers the
    pictures it adds from 1 each time, which collides with the shape ids."""
    tag = qn("wp:docPr")
    for part in [doc.element.body] + [h._element for s_ in doc.sections
                                      for h in (s_.header, s_.footer)]:
        for i, el in enumerate(part.iter(tag), 1):
            el.set("id", str(i))


def main(out="JBM_DMS_Service_Mobile_App_User_Manual.docx"):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.0

    page_setup(doc)
    cover(doc)
    front_matter(doc)
    introduction(doc)
    steps(doc)
    back_matter(doc)
    renumber_drawings(doc)
    doc.save(out)
    print("wrote", out, os.path.getsize(out) // 1024, "KB")


if __name__ == "__main__":
    main(*sys.argv[1:])
